from fastapi import APIRouter, Depends, HTTPException
from sqlalchemy.orm import Session
from sqlalchemy import func
from pydantic import BaseModel
from typing import Optional
from app.database.session import get_db
from app.models.user import User, UserRole
from app.models.patient import Patient
from app.auth.deps import require_role, get_current_user
from app.core.config import settings
import uuid
from datetime import datetime, date

router = APIRouter()

@router.get("/me")
def get_my_patient_profile(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    if current_user.role != UserRole.PATIENT:
        raise HTTPException(status_code=400, detail="Only patient accounts can access this profile endpoint")
    
    patient = db.query(Patient).filter(Patient.user_id == current_user.id).first()
    if not patient:
        patient = db.query(Patient).filter(func.lower(Patient.email) == current_user.username.lower().strip()).first()
        
    if not patient:
        raise HTTPException(status_code=404, detail="Patient clinical profile not found")
        
    return {
        "id": patient.id,
        "patient_access_id": patient.patient_access_id,
        "first_name": patient.first_name,
        "last_name": patient.last_name,
        "name": f"{patient.first_name} {patient.last_name}",
        "email": patient.email,
        "phone": patient.phone,
        "date_of_birth": str(patient.date_of_birth) if patient.date_of_birth else None,
        "diabetes_type": patient.diabetes_type or "Not specified",
        "year_of_diagnosis": patient.year_of_diagnosis,
        "existing_eye_conditions": patient.existing_eye_conditions,
        "clinic_name": "Retinal Care Unit & Diabetic Eye Clinic",
        "attending_physician": "Dr. Screening",
    }


class CreatePatientRequest(BaseModel):
    first_name: str
    last_name: str
    email: str  # Required, used as login identifier
    phone: Optional[str] = None
    date_of_birth: str  # YYYY-MM-DD
    diabetes_type: str  # "Type 1", "Type 2", "Gestational", "Pre-diabetic", "Not diabetic"
    year_of_diagnosis: Optional[int] = None
    existing_eye_conditions: Optional[str] = None


@router.post("/", status_code=201)
def create_patient(
    data: CreatePatientRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_role([UserRole.HEALTHCARE_WORKER, UserRole.ADMIN]))
):
    email_clean = data.email.lower().strip()
    if not email_clean:
        raise HTTPException(status_code=400, detail="Email address is required.")

    # Duplicate check on email
    existing_patient = db.query(Patient).filter(func.lower(Patient.email) == email_clean).first()
    if existing_patient:
        raise HTTPException(
            status_code=400,
            detail=f"A patient with this email already exists (Patient ID: {existing_patient.patient_access_id})."
        )

    existing_user = db.query(User).filter(func.lower(User.username) == email_clean).first()
    if existing_user:
        linked_patient = db.query(Patient).filter(Patient.user_id == existing_user.id).first()
        pat_id_str = f" (Patient ID: {linked_patient.patient_access_id})" if linked_patient else ""
        raise HTTPException(
            status_code=400,
            detail=f"A patient with this email already exists{pat_id_str}."
        )

    # Parse date_of_birth
    dob_obj = None
    if data.date_of_birth:
        try:
            dob_obj = datetime.strptime(data.date_of_birth.strip(), "%Y-%m-%d").date()
        except ValueError:
            raise HTTPException(status_code=400, detail="Invalid Date of Birth format. Please use YYYY-MM-DD.")
    else:
        raise HTTPException(status_code=400, detail="Date of Birth is required for patient verification.")

    # Create linked user account with hashed_password = NULL and is_activated = False
    user = User(
        username=email_clean,
        hashed_password=None,
        role=UserRole.PATIENT,
        is_active=True,
        is_activated=False,
        require_password_change=True,
    )
    db.add(user)
    db.commit()
    db.refresh(user)

    # Create patient profile with full clinical biodata
    patient_access_id = f"MV-PAT-{uuid.uuid4().hex[:6].upper()}"
    patient = Patient(
        user_id=user.id,
        patient_access_id=patient_access_id,
        first_name=data.first_name.strip(),
        last_name=data.last_name.strip(),
        date_of_birth=dob_obj,
        diabetes_type=data.diabetes_type.strip() if data.diabetes_type else "Not diabetic",
        year_of_diagnosis=data.year_of_diagnosis,
        existing_eye_conditions=data.existing_eye_conditions.strip() if data.existing_eye_conditions else None,
        email=email_clean,
        phone=data.phone.strip() if data.phone else None,
    )
    db.add(patient)
    db.commit()
    db.refresh(patient)

    # TODO: Automated email delivery hook.
    # When email infrastructure is provisioned, send a welcome/activation email to patient.email.
    # For now, login details (email & Patient ID) are handed off out-of-band by the clinician.

    return {
        "patient_id": patient.id,
        "patient_access_id": patient_access_id,
        "first_name": patient.first_name,
        "last_name": patient.last_name,
        "name": f"{patient.first_name} {patient.last_name}",
        "email": patient.email,
        "date_of_birth": str(patient.date_of_birth) if patient.date_of_birth else None,
        "diabetes_type": patient.diabetes_type,
        "year_of_diagnosis": patient.year_of_diagnosis,
        "existing_eye_conditions": patient.existing_eye_conditions,
        "account_status": "PENDING_ACTIVATION",
        "message": "Patient registered successfully. Share login details with the patient.",
    }


@router.get("/")
def list_patients(
    db: Session = Depends(get_db),
    current_user: User = Depends(require_role([UserRole.HEALTHCARE_WORKER, UserRole.ADMIN]))
):
    patients = db.query(Patient).all()
    return [
        {
            "id": p.id,
            "patient_access_id": p.patient_access_id,
            "name": f"{p.first_name} {p.last_name}",
            "first_name": p.first_name,
            "last_name": p.last_name,
            "email": p.email,
            "phone": p.phone,
            "date_of_birth": str(p.date_of_birth) if p.date_of_birth else None,
            "diabetes_type": p.diabetes_type,
            "year_of_diagnosis": p.year_of_diagnosis,
            "existing_eye_conditions": p.existing_eye_conditions,
            "username": p.user.username if p.user else None,
            "is_activated": (p.user.hashed_password is not None and getattr(p.user, 'is_activated', True)) if p.user else False,
            "account_status": "ACTIVE" if (p.user and p.user.hashed_password is not None) else "PENDING_ACTIVATION",
            "portal_active": p.user.hashed_password is not None if p.user else False,
        }
        for p in patients
    ]


@router.delete("/{patient_id}")
def delete_patient(
    patient_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_role([UserRole.HEALTHCARE_WORKER, UserRole.ADMIN]))
):
    patient = db.query(Patient).filter(Patient.id == patient_id).first()
    if not patient:
        raise HTTPException(status_code=404, detail="Patient not found")

    user_id = patient.user_id

    # Delete associated screenings and reports
    from app.models.screening import Screening
    from app.models.report import Report
    from app.rag.store import chroma_store

    screenings = db.query(Screening).filter(Screening.patient_id == patient_id).all()
    screening_ids = [s.id for s in screenings]

    if screening_ids:
        db.query(Report).filter(Report.screening_id.in_(screening_ids)).delete(synchronize_session=False)
        db.query(Screening).filter(Screening.id.in_(screening_ids)).delete(synchronize_session=False)

    # Delete patient vector embeddings from Chroma
    chroma_store.delete_patient_reports(patient_id)

    # Delete patient profile
    db.delete(patient)

    # Delete associated user account
    if user_id:
        user = db.query(User).filter(User.id == user_id).first()
        if user:
            db.delete(user)

    db.commit()
    return {"message": "Patient and all associated records deleted successfully"}


@router.get("/export/excel")
@router.get("/export/csv")
def export_patients_excel(
    db: Session = Depends(get_db),
    current_user: User = Depends(require_role([UserRole.HEALTHCARE_WORKER, UserRole.ADMIN]))
):
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from fastapi.responses import Response
    from app.models.screening import Screening

    wb = Workbook()

    # Shared Styles
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    fill_header = PatternFill(start_color="C85A32", end_color="C85A32", fill_type="solid")
    align_header = Alignment(horizontal="center", vertical="center", wrap_text=True)

    border_thin = Side(border_style="thin", color="D1D5DB")
    box_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)

    fill_dr = PatternFill(start_color="FFE4E6", end_color="FFE4E6", fill_type="solid")
    font_dr = Font(name="Calibri", size=11, bold=True, color="991B1B")

    fill_nodr = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
    font_nodr = Font(name="Calibri", size=11, bold=True, color="065F46")

    # =========================================================================
    # SHEET 1: ALL SCREENING OUTPUTS (PRIMARY SHEET)
    # =========================================================================
    ws1 = wb.active
    ws1.title = "All Screening Outputs"
    ws1.views.sheetView[0].showGridLines = True

    headers_screenings = [
        "Screening Reference",
        "Date & Time (UTC)",
        "Patient Access Code",
        "Patient Name",
        "Patient Email",
        "Diagnostic Finding",
        "Model Confidence",
        "DR Risk Score",
        "Normal Score",
        "Assessed Risk Level",
        "Grok AI Clinical Context Analysis",
        "Recommended Clinical Management",
        "Retinal Scan Image Path",
        "Grad-CAM Heatmap Image Path"
    ]

    ws1.append(headers_screenings)
    ws1.row_dimensions[1].height = 28

    for col_num in range(1, len(headers_screenings) + 1):
        cell = ws1.cell(row=1, column=col_num)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_header
        cell.border = box_border

    # Fetch ALL screening outputs, newest first
    all_screenings = db.query(Screening).order_by(Screening.created_at.desc()).all()

    for s in all_screenings:
        patient = s.patient
        patient_name = f"{patient.first_name} {patient.last_name}" if patient else "N/A"
        patient_code = patient.patient_access_id if patient else "N/A"
        patient_email = patient.email if patient else "N/A"
        date_str = s.created_at.strftime('%Y-%m-%d %H:%M:%S') if s.created_at else "N/A"
        
        prob_dr = getattr(s, 'probability_dr', s.confidence if s.prediction == "DR PRESENT" else 1 - s.confidence) or 0.0
        prob_no_dr = getattr(s, 'probability_no_dr', 1 - s.confidence if s.prediction == "DR PRESENT" else s.confidence) or 0.0
        ai_ctx = getattr(s, 'ai_context', '') or ""

        row_data = [
            s.screening_id,
            date_str,
            patient_code,
            patient_name,
            patient_email,
            s.prediction,
            f"{s.confidence * 100:.1f}%",
            f"{prob_dr * 100:.1f}%",
            f"{prob_no_dr * 100:.1f}%",
            s.risk_level,
            ai_ctx,
            s.recommendation,
            s.image_path or "",
            s.heatmap_path or ""
        ]

        ws1.append(row_data)
        current_row = ws1.max_row
        ws1.row_dimensions[current_row].height = 24

        # Highlight Finding Column (Col 6)
        finding_cell = ws1.cell(row=current_row, column=6)
        if s.prediction == "DR PRESENT":
            finding_cell.fill = fill_dr
            finding_cell.font = font_dr
        elif s.prediction == "NO DR":
            finding_cell.fill = fill_nodr
            finding_cell.font = font_nodr

        for col_num in range(1, len(headers_screenings) + 1):
            cell = ws1.cell(row=current_row, column=col_num)
            cell.border = box_border
            if col_num in [11, 12]:  # Wrap text for long Grok context & recommendation
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            elif col_num in [1, 2, 3, 6, 7, 8, 9, 10]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    for col in ws1.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        # Give long text columns comfortable fixed width
        if col_letter in ['K', 'L']:
            ws1.column_dimensions[col_letter].width = 45
        elif col_letter in ['M', 'N']:
            ws1.column_dimensions[col_letter].width = 35
        else:
            ws1.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 40)

    # =========================================================================
    # SHEET 2: PATIENT DIRECTORY SUMMARY
    # =========================================================================
    ws2 = wb.create_sheet(title="Patient Directory Summary")
    ws2.views.sheetView[0].showGridLines = True

    headers_patients = [
        "Patient ID",
        "Patient Access Code",
        "Full Name",
        "Email Address",
        "Phone Number",
        "Portal Username",
        "Account Status",
        "Total Screenings Recorded",
        "Latest Screening ID",
        "Latest Finding",
        "Latest Confidence",
        "Latest Risk Level"
    ]

    ws2.append(headers_patients)
    ws2.row_dimensions[1].height = 28

    for col_num in range(1, len(headers_patients) + 1):
        cell = ws2.cell(row=1, column=col_num)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_header
        cell.border = box_border

    patients = db.query(Patient).all()
    for p in patients:
        p_screenings = db.query(Screening).filter(Screening.patient_id == p.id).order_by(Screening.created_at.desc()).all()
        latest = p_screenings[0] if p_screenings else None

        row_data_p = [
            p.id,
            p.patient_access_id,
            f"{p.first_name} {p.last_name}",
            p.email or "",
            p.phone or "",
            p.user.username if p.user else "",
            "Active" if (p.user and p.user.hashed_password) else "Pending Activation",
            len(p_screenings),
            latest.screening_id if latest else "N/A",
            latest.prediction if latest else "N/A",
            f"{latest.confidence * 100:.1f}%" if latest else "N/A",
            latest.risk_level if latest else "N/A"
        ]

        ws2.append(row_data_p)
        current_row_2 = ws2.max_row
        ws2.row_dimensions[current_row_2].height = 22

        finding_cell_2 = ws2.cell(row=current_row_2, column=10)
        if latest and latest.prediction == "DR PRESENT":
            finding_cell_2.fill = fill_dr
            finding_cell_2.font = font_dr
        elif latest and latest.prediction == "NO DR":
            finding_cell_2.fill = fill_nodr
            finding_cell_2.font = font_nodr

        for col_num in range(1, len(headers_patients) + 1):
            cell = ws2.cell(row=current_row_2, column=col_num)
            cell.border = box_border
            if col_num in [1, 2, 7, 8, 9, 11, 12]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    for col in ws2.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws2.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 40)

    # Save to BytesIO stream
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=MedVisionAI_All_Screening_Outputs.xlsx"}
    )


